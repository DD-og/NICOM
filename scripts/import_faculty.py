"""Turn the institute's faculty list (.xlsx) into portal accounts.

Run: python3 scripts/import_faculty.py IMNU_Faculty_List.xlsx
The list needs a header row with the faculty name ("Teaching Staff ..."), "Area", "Email ID", "Dept" (PG/UG)
and "All Area": the conference tracks that person may review, separated by ";". Every track must be one of
DOMAINS in lib/config.ts; a person reviews papers in those tracks only.

Writes (all git-ignored, they hold password data):
  data/faculty-import.csv              rows for the Sheet's Faculty tab (password hashes only)
  data/faculty-logins.local.xlsx       name, email, tracks, login ID and password per person, to hand out
  data/faculty-credentials.local.json  passwords, reused on re-runs so logins stay the same
People already in data/faculty-import.csv keep their Faculty ID and password; the admin account is kept from
data/seed.json.
"""
import csv, hashlib, json, os, re, secrets, sys
from pathlib import Path

import openpyxl

MAX_PAPERS = 8
FIRST_ID = 101  # FAC101 onwards, so nothing left over from the demo accounts (FAC01-FAC10) can attach to a real person
IMPORT_FILE = Path("data/faculty-import.csv")
CRED_FILE = Path("data/faculty-credentials.local.json")
ALPHABET = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # no look-alikes (0/O, 1/l/I)
HEADERS = ["Faculty ID", "Username", "Name", "Email", "Role", "Domains", "Max Papers", "Active", "Password Hash", "Area", "Programme"]
TRACKS = re.findall(r'"([^"]+)"', re.search(r"DOMAINS = \[(.*?)\];", Path("lib/config.ts").read_text(), re.S).group(1))


def new_password() -> str:
    return "-".join("".join(secrets.choice(ALPHABET) for _ in range(4)) for _ in range(3))


def hash_pw(pw: str) -> str:
    # Must match lib/password.ts (scrypt N=16384, r=8, p=1, 64 bytes)
    salt = os.urandom(16)
    dk = hashlib.scrypt(pw.encode(), salt=salt, n=16384, r=8, p=1, dklen=64)
    return f"scrypt${salt.hex()}${dk.hex()}"


def tidy(s: str) -> str:
    return re.sub(r"\s*&\s*", " & ", re.sub(r"\s+", " ", s.strip()))


def read_people(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb:
        cols = None
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if cols is None:
                lower = [c.lower() for c in cells]
                if any(c.startswith("email") for c in lower):
                    find = lambda test: next((i for i, c in enumerate(lower) if test(c)), None)
                    cols = {
                        "name": find(lambda c: "staff" in c or c == "name"),
                        "area": find(lambda c: c == "area"),
                        "dept": find(lambda c: c in ("dept", "programme", "program")),
                        "email": find(lambda c: c.startswith("email")),
                        "tracks": find(lambda c: c in ("all area", "all areas", "tracks")),
                    }
                    if cols["tracks"] is None:
                        sys.exit(f"Sheet {ws.title!r} has no 'All Area' column with each person's tracks.")
                continue
            if cells[cols["email"]] and cells[cols["name"]]:
                dept = cells[cols["dept"]].upper() if cols["dept"] is not None else ""
                yield {
                    "name": cells[cols["name"]],
                    "area": tidy(cells[cols["area"]]) if cols["area"] is not None else "",
                    "programme": dept or ("UG" if "UG" in ws.title.upper() else "PG"),
                    "email": cells[cols["email"]].lower(),
                    "tracks": [tidy(t) for t in cells[cols["tracks"]].split(";") if t.strip()],
                }


def main(src: str):
    creds = json.loads(CRED_FILE.read_text()) if CRED_FILE.exists() else {}
    previous = {r["Email"].lower(): r for r in csv.DictReader(IMPORT_FILE.open())} if IMPORT_FILE.exists() else {}
    next_id = max([int(re.sub(r"\D", "", r["Faculty ID"]) or 0) for r in previous.values() if r["Role"] == "faculty"] + [FIRST_ID - 1]) + 1
    seed = json.loads(Path("data/seed.json").read_text())
    admin = next(f for f in seed["faculty"] if f["role"] == "master")
    rows = [[admin["id"], admin["username"], admin["name"], admin["email"], "master", "", 0, "TRUE", admin["passwordHash"], "", ""]]
    logins, seen, problems = [], set(), []
    known = {t.lower(): t for t in TRACKS}
    for p in read_people(src):
        user = p["email"].split("@")[0]
        if user in seen:
            sys.exit(f"Two people would get the login ID {user!r}; fix the emails first.")
        seen.add(user)
        unknown = [t for t in p["tracks"] if t.lower() not in known]
        if unknown or not p["tracks"]:
            sys.exit(f"{p['name']}: tracks {unknown or '(none)'} are not conference tracks in lib/config.ts.")
        tracks = [known[t.lower()] for t in p["tracks"]]
        if not p["email"].endswith("@nirmauni.ac.in"):
            problems.append(f"{p['name']}: email {p['email']} is not @nirmauni.ac.in")
        creds.setdefault(user, new_password())
        old = previous.get(p["email"])
        if old and old["Username"] == user:
            fid, pw_hash = old["Faculty ID"], old["Password Hash"]
        else:
            fid, pw_hash = f"FAC{next_id}", hash_pw(creds[user])
            next_id += 1
        rows.append([fid, user, p["name"], p["email"], "faculty", "; ".join(tracks), MAX_PAPERS, "TRUE", pw_hash, p["area"], p["programme"]])
        logins.append([p["programme"], p["name"], p["area"], p["email"], "; ".join(tracks), user, creds[user]])
    CRED_FILE.write_text(json.dumps(creds, indent=2))

    with IMPORT_FILE.open("w", newline="") as f:
        csv.writer(f).writerows([HEADERS, *rows])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logins"
    ws.append(["Programme", "Name", "Area", "Email", "Review tracks", "Login ID", "Password"])
    admin_pw = json.loads(Path("data/credentials.local.json").read_text()).get("admin", "(unchanged)") if Path("data/credentials.local.json").exists() else "(unchanged)"
    ws.append(["", admin["name"] + " (admin)", "", admin["email"], "All", admin["username"], admin_pw])
    for r in logins:
        ws.append(r)
    for col, width in zip("ABCDEFG", [11, 30, 26, 36, 90, 22, 18]):
        ws.column_dimensions[col].width = width
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save("data/faculty-logins.local.xlsx")

    print(f"{len(logins)} faculty + 1 admin -> {IMPORT_FILE}, data/faculty-logins.local.xlsx")
    for msg in problems:
        print("note:", msg)


if __name__ == "__main__":
    main(sys.argv[1])
