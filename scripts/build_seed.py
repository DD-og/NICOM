"""Combine dummy_source.json + drive_ids.json into the app seed and a multi-tab .xlsx
that is uploaded to Google Drive (auto-converted to a Google Sheet).

Run: python3 scripts/build_seed.py [xlsx_out_path]
Test logins: one ID + password per account. Passwords are random, kept in data/credentials.local.json
(git-ignored, created on first run, reused afterwards); only their scrypt hashes go into seed.json.
"""
import hashlib, json, os, random, secrets, sys
from datetime import datetime, timedelta
from pathlib import Path

random.seed(7)
src = json.loads(Path("data/dummy_source.json").read_text())
drive = json.loads(Path("data/drive_ids.json").read_text())


CRED_FILE = Path("data/credentials.local.json")
ALPHABET = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # no look-alikes (0/O, 1/l/I)


def new_password() -> str:
    return "-".join("".join(secrets.choice(ALPHABET) for _ in range(4)) for _ in range(3))


def hash_pw(pw: str) -> str:
    # Must match lib/password.ts (scrypt N=16384, r=8, p=1, 64 bytes)
    salt = os.urandom(16)
    dk = hashlib.scrypt(pw.encode(), salt=salt, n=16384, r=8, p=1, dklen=64)
    return f"scrypt${salt.hex()}${dk.hex()}"


def username(f) -> str:
    return "admin" if f["role"] == "master" else f["email"].split("@")[0]


creds = json.loads(CRED_FILE.read_text()) if CRED_FILE.exists() else {}
faculty = []
for f in src["faculty"]:
    u = username(f)
    creds.setdefault(u, new_password())
    faculty.append({**f, "username": u, "active": True, "maxLoad": 8 if f["role"] == "faculty" else 0,
                    "passwordHash": hash_pw(creds[u])})
CRED_FILE.write_text(json.dumps(creds, indent=2))

papers = []
for p in src["papers"]:
    papers.append({**p, "driveLink": f"https://docs.google.com/document/d/{drive['papers'][p['id']]}/view"})

reviewers = [f for f in faculty if f["role"] == "faculty"]
load = {f["id"]: 0 for f in reviewers}
assignments, reviews, decisions = [], [], []
# Reviews are drawn from five demo outcomes (kept so the random sequence, and the demo data, stay the same)
# and stored in the committee's evaluation matrix: nine 1-5 statements, the overall evaluation and three
# internal Yes/No nominations.
RECS = ["Accept", "Accept with Minor Revisions", "Major Revisions", "Reject", "Recommended for Journal Publication"]
OVERALL = {
    "Accept": "Accept as Submitted with Minor Formatting changes if required",
    "Accept with Minor Revisions": "Accept With Suggested Minor Revisions",
    "Major Revisions": "Accept After Major Revision",
    "Reject": "Reject",
    "Recommended for Journal Publication": "Accept as Submitted with Minor Formatting changes if required",
}
COMMENTS = {
    "Accept": "The paper is well structured with a clear research question and a sound methodology. The contribution to the Indian context is evident.",
    "Accept with Minor Revisions": "A relevant study with an adequate design. Please tighten the literature review, report robustness checks and fix minor formatting issues.",
    "Major Revisions": "The topic is interesting, but the methodology section needs substantial work: justify the sample, address endogeneity and expand the discussion of limitations.",
    "Reject": "The paper lacks a clear theoretical framing and the empirical analysis does not support the stated conclusions. The contribution is limited in its current form.",
    "Recommended for Journal Publication": "An excellent, rigorous paper with a strong contribution and clear implications for practice.",
}


def matrix(rec: str, s: int, sub: dict) -> dict:
    """Spread the demo sub-scores over the nine statements and derive the nominations."""
    o, m, r, c = sub["originality"], sub["methodology"], sub["relevance"], sub["clarity"]
    return {
        "novelty": o, "literature": o, "methodology": m, "analysis": m, "organisation": c,
        "contribution": s, "relevance": r, "implications": s, "usefulness": r,
        "recommendation": OVERALL[rec],
        "journal": "Yes" if rec == "Recommended for Journal Publication" else "No",
        "scopus": "Yes" if s >= 4 else "No",
        "bestPaper": "Yes" if rec == "Recommended for Journal Publication" else "No",
    }

base = datetime(2026, 9, 26, 10, 0)
for idx, p in enumerate(papers, start=1):
    if idx > 18:  # 19-24 left unassigned for the master admin to distribute
        continue
    n_rev = 2  # conference policy: two reviewers per paper
    pool = sorted([f for f in reviewers if p["domain"] in f["domains"]], key=lambda f: (load[f["id"]], f["id"]))
    for f in pool[:n_rev]:
        load[f["id"]] += 1
        a_id = f"ASG-{len(assignments) + 1:03d}"
        assigned_at = base + timedelta(hours=idx)
        done = idx <= 10
        assignments.append({
            "id": a_id, "paperId": p["id"], "facultyId": f["id"],
            "assignedAt": assigned_at.isoformat(), "dueDate": (assigned_at + timedelta(days=14)).date().isoformat(),
            "status": "Completed" if done else "Pending", "assignedBy": "MASTER",
        })
        if done:
            rec = random.choice(RECS)
            s = {"Accept": 4, "Accept with Minor Revisions": 4, "Major Revisions": 3, "Reject": 2, "Recommended for Journal Publication": 5}[rec]
            sub = {k: max(1, min(5, s + random.choice([-1, 0, 0, 1]))) for k in ["originality", "methodology", "relevance", "clarity"]}
            reviews.append({
                "id": f"REV-{len(reviews) + 1:03d}", "assignmentId": a_id, "paperId": p["id"], "facultyId": f["id"],
                **matrix(rec, s, sub),
                "commentsToAuthors": COMMENTS[rec],
                "confidentialComments": "" if random.random() < 0.6 else "Check for overlap with an earlier working paper by the same group.",
                "submittedAt": (assigned_at + timedelta(days=random.randint(2, 9))).isoformat(),
            })

# Paper 10: both reviewers agree, so the portal has auto-drafted the decision (shows the auto-draft flow).
r10 = [r for r in reviews if r["paperId"] == papers[9]["id"]]
for k in ["novelty", "literature", "methodology", "analysis", "organisation", "contribution", "relevance",
          "implications", "usefulness", "recommendation", "journal", "scopus", "bestPaper", "commentsToAuthors"]:
    r10[1][k] = r10[0][k]
r10[1]["commentsToAuthors"] = "Strong empirical work with clear managerial implications. Minor language edits would improve readability."
decisions.append({
    "paperId": papers[9]["id"], "finalDecision": r10[0]["recommendation"], "decidedBy": "AUTO",
    "decidedAt": max(r["submittedAt"] for r in r10), "emailStatus": "Draft", "emailSentAt": "",
    "notes": "Auto-drafted from reviewer consensus",
})

# Final decisions for the first four papers (made by the master admin); two of the emails already sent.
for i, p in enumerate(papers[:4]):
    revs = [r for r in reviews if r["paperId"] == p["id"]]
    decisions.append({
        "paperId": p["id"], "finalDecision": revs[0]["recommendation"], "decidedBy": "MASTER",
        "decidedAt": (base + timedelta(days=12, hours=i)).isoformat(),
        "emailStatus": "Sent" if i < 2 else "Draft",
        "emailSentAt": (base + timedelta(days=12, hours=i + 2)).isoformat() if i < 2 else "",
        "notes": "",
    })

seed = {"domains": src["domains"], "faculty": faculty, "papers": papers,
        "assignments": assignments, "reviews": reviews, "decisions": decisions,
        "settings": {"driveFolder": f"https://drive.google.com/drive/folders/{drive['submissionsFolder']}"}}
Path("data/seed.json").write_text(json.dumps(seed, indent=2))
print(f"seed: {len(papers)} papers, {len(assignments)} assignments, {len(reviews)} reviews, {len(decisions)} decisions")

# ---------- XLSX (becomes the Google Sheet) ----------
if len(sys.argv) > 1:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    sys.path.insert(0, str(Path(__file__).parent))
    cols = json.loads(Path("lib/sheet-columns.json").read_text())
    wb = Workbook(); wb.remove(wb.active)
    tables = {"Submissions": papers, "Faculty": faculty, "Assignments": assignments, "Reviews": reviews, "Decisions": decisions}
    for tab, rows in tables.items():
        ws = wb.create_sheet(tab)
        spec = cols[tab]
        ws.append([c["header"] for c in spec])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3A68")
        for r in rows:
            out = []
            for c in spec:
                v = r.get(c["key"], "")
                if isinstance(v, list): v = ", ".join(v)
                if isinstance(v, bool): v = "TRUE" if v else "FALSE"
                out.append(v)
            ws.append(out)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(12, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(sys.argv[1])
    print("xlsx:", sys.argv[1])
